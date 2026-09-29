#!/usr/bin/env python3
"""Slide Renderer - 将 PPTX 幻灯片渲染为 PNG 图片。

用于 VLM 处理的图片输入。
"""

import io
import subprocess
import tempfile
from pathlib import Path


class SlideRenderer:
    """将 PPTX 幻灯片渲染为 PNG 图片。"""

    def __init__(self, pptx_path: str, output_dir: str | None = None):
        self.pptx_path = Path(pptx_path)
        self.output_dir = Path(output_dir) if output_dir else None

    def render_all_slides(self) -> list[Path]:
        """渲染所有幻灯片为 PNG。

        优先使用 LibreOffice，失败则回退到其他方式。
        """
        images = []

        # 尝试 LibreOffice
        if self._has_libreoffice():
            images = self._render_with_libreoffice()
            if images:
                return images

        # 回退：使用 pdf2image（如果可用）
        images = self._render_with_pdf2image()
        if images:
            return images

        # 最后回退：用 PIL 按元素坐标重绘（不依赖外部工具）
        return self._render_with_pil()

    def _has_libreoffice(self) -> bool:
        """检查 LibreOffice 是否可用。"""
        try:
            subprocess.run(
                ["soffice", "--version"],
                capture_output=True,
                timeout=5,
            )
            return True
        except (FileNotFoundError, subprocess.TimeoutExpired):
            return False

    def _render_with_libreoffice(self) -> list[Path]:
        """使用 LibreOffice 渲染。"""
        output_dir = self.output_dir or Path(tempfile.mkdtemp(prefix="pptx_render_"))
        output_dir.mkdir(parents=True, exist_ok=True)

        try:
            # 先转 PDF
            subprocess.run(
                [
                    "soffice",
                    "--headless",
                    "--convert-to", "pdf",
                    "--outdir", str(output_dir),
                    str(self.pptx_path),
                ],
                capture_output=True,
                timeout=60,
            )

            pdf_path = output_dir / f"{self.pptx_path.stem}.pdf"
            if not pdf_path.exists():
                return []

            # PDF 转 PNG
            return self._pdf_to_png(pdf_path, output_dir)

        except (subprocess.TimeoutExpired, Exception):
            return []

    def _pdf_to_png(self, pdf_path: Path, output_dir: Path) -> list[Path]:
        """将 PDF 转为 PNG 序列。"""
        try:
            from pdf2image import convert_from_path
            images_pil = convert_from_path(str(pdf_path), dpi=150)
            png_paths = []
            for i, img in enumerate(images_pil):
                png_path = output_dir / f"slide_{i:03d}.png"
                img.save(png_path, "PNG")
                png_paths.append(png_path)
            return png_paths
        except ImportError:
            return []

    def _render_with_pdf2image(self) -> list[Path]:
        """直接用 pdf2image（需要已有 PDF）。"""
        # 这个方法需要先有 PDF，暂时跳过
        return []

    def _extract_thumbnails(self) -> list[Path]:
        """提取 PPTX 内嵌的缩略图（最后回退方案）。"""
        # python-pptx 不直接支持，返回空
        return []

    def _render_with_pil(self, dpi: int = 150) -> list[Path]:
        """用 PIL 按元素坐标重绘每页，作为 VLM 的视觉输入。

        不依赖 LibreOffice 等外部工具，能反映页面版面结构
        （文本/图片/表格的位置与相对大小），足以让 VLM 判断阅读顺序。
        """
        try:
            from pptx import Presentation
            from pptx.util import Emu
            from PIL import Image, ImageDraw
        except ImportError:
            return []

        output_dir = self.output_dir or Path(tempfile.mkdtemp(prefix="pptx_render_"))
        output_dir.mkdir(parents=True, exist_ok=True)

        try:
            prs = Presentation(str(self.pptx_path))
        except Exception:
            return []

        # 幻灯片尺寸（EMU）
        slide_w = prs.slide_width
        slide_h = prs.slide_height
        if not slide_w or not slide_h:
            slide_w, slide_h = 9144000, 6858000  # 默认 10x7.5 inch

        # 像素尺寸（按 dpi 缩放）
        px_w = int(Emu(slide_w).inches * dpi)
        px_h = int(Emu(slide_h).inches * dpi)
        # 限制最大边长，避免图片过大
        max_side = 2000
        scale = 1.0
        if max(px_w, px_h) > max_side:
            scale = max_side / max(px_w, px_h)
            px_w = int(px_w * scale)
            px_h = int(px_h * scale)

        png_paths = []
        for idx, slide in enumerate(prs.slides):
            img = Image.new("RGB", (px_w, px_h), "white")
            draw = ImageDraw.Draw(img)

            for shape in slide.shapes:
                self._draw_shape(img, draw, shape, scale)

            png_path = output_dir / f"slide_{idx:03d}.png"
            img.save(png_path, "PNG")
            png_paths.append(png_path)

        return png_paths

    @staticmethod
    def _emu_to_px(emu, scale: float) -> int:
        """EMU → 像素（基准 dpi=150，再乘缩放比例）。"""
        return int(emu / 914400 * 150 * scale)

    def _draw_shape(self, img, draw, shape, scale: float):
        """在 PIL 画布上绘制单个 shape。"""
        import io
        from PIL import Image

        if shape.left is None or shape.top is None:
            return

        x1 = self._emu_to_px(shape.left, scale)
        y1 = self._emu_to_px(shape.top, scale)
        w = self._emu_to_px(shape.width or 0, scale)
        h = self._emu_to_px(shape.height or 0, scale)

        if w <= 0 or h <= 0:
            return

        # 图片：直接粘贴到画布
        if shape.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE
            try:
                blob = shape.image.blob
                pil_img = Image.open(io.BytesIO(blob)).convert("RGB")
                pil_img = pil_img.resize((max(w, 1), max(h, 1)))
                img.paste(pil_img, (x1, y1))
            except Exception:
                draw.rectangle([x1, y1, x1 + w, y1 + h], outline=(0, 0, 255), width=2)
            return

        # 表格：画格子 + 文本
        if getattr(shape, "has_table", False):
            self._draw_table(draw, shape.table, x1, y1, w, h)
            return

        # 文本
        if shape.has_text_frame:
            text = shape.text_frame.text
            draw.rectangle([x1, y1, x1 + w, y1 + h], outline=(0, 0, 0), width=1)
            if text.strip():
                self._draw_text(draw, text, x1, y1, w, h)
            return

        # 其他形状：画矩形
        draw.rectangle([x1, y1, x1 + w, y1 + h], outline=(0, 0, 0), width=1)

    def _draw_table(self, draw, table, x1, y1, w, h):
        """绘制表格网格与文本。"""
        nrows = len(table.rows)
        ncols = len(table.columns)
        if nrows == 0 or ncols == 0:
            draw.rectangle([x1, y1, x1 + w, y1 + h], outline=(0, 0, 0), width=2)
            return

        row_h = h / nrows
        col_w = w / ncols
        for r in range(nrows):
            for c in range(ncols):
                cx1 = x1 + int(c * col_w)
                cy1 = y1 + int(r * row_h)
                cx2 = x1 + int((c + 1) * col_w)
                cy2 = y1 + int((r + 1) * row_h)
                draw.rectangle([cx1, cy1, cx2, cy2], outline=(0, 0, 0), width=1)
                try:
                    cell_text = table.cell(r, c).text
                except Exception:
                    cell_text = ""
                if cell_text.strip():
                    self._draw_text(draw, cell_text, cx1 + 2, cy1 + 2, cx2 - cx1 - 4, cy2 - cy1 - 4)

    def _draw_text(self, draw, text: str, x1: int, y1: int, w: int, h: int):
        """在限定区域内绘制文本（按可用空间截断）。"""
        from PIL import ImageFont

        # 根据区域高度估算字号
        font_size = max(10, min(int(h * 0.7), 32))
        try:
            font = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", font_size)
        except Exception:
            font = ImageFont.load_default()

        lines = text.split("\n")
        line_h = font_size + 2
        y = y1
        for line in lines:
            if y + line_h > y1 + h:
                break
            draw.text((x1 + 2, y), line, fill=(0, 0, 0), font=font)
            y += line_h

    def render_single_slide(self, slide_index: int) -> Path | None:
        """渲染单张幻灯片。"""
        all_images = self.render_all_slides()
        if slide_index < len(all_images):
            return all_images[slide_index]
        return None
