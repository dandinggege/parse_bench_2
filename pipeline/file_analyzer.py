#!/usr/bin/env python3
"""File Analyzer - PPTX 文件分析器。

双路解析策略：
1. Native Parsing - 用 python-pptx 提取文本/表格/图表的结构化数据
2. Slide Rendering - 将幻灯片渲染为 PNG 图片供 VLM 处理

融合规则：Native 优先，VLM 补充 Native 解不出的部分（如嵌入图片、SmartArt）
"""

import io
from dataclasses import dataclass, field
from typing import Optional

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu


@dataclass
class TextElement:
    """文本元素。"""
    text: str
    bbox: tuple[int, int, int, int]  # x1, y1, x2, y2 (EMU)
    font_size: float = 0
    is_title: bool = False
    shape_type: str = ""


@dataclass
class TableElement:
    """表格元素。"""
    rows: int
    cols: int
    cells: list[list[str]]
    bbox: tuple[int, int, int, int]
    html: str = ""


@dataclass
class ChartElement:
    """图表元素。"""
    chart_type: str
    data: dict
    bbox: tuple[int, int, int, int]
    html: str = ""


@dataclass
class ImageElement:
    """图片元素（需要 VLM 处理）。"""
    image_bytes: bytes
    bbox: tuple[int, int, int, int]
    shape_type: str = ""


@dataclass
class ShapeElement:
    """形状元素（流程图等）。"""
    shape_type: str
    text: str
    bbox: tuple[int, int, int, int]
    fill_color: str = ""


@dataclass
class SlideData:
    """单页幻灯片数据。"""
    slide_index: int
    title: str = ""
    texts: list[TextElement] = field(default_factory=list)
    tables: list[TableElement] = field(default_factory=list)
    charts: list[ChartElement] = field(default_factory=list)
    images: list[ImageElement] = field(default_factory=list)
    shapes: list[ShapeElement] = field(default_factory=list)
    width: int = 0
    height: int = 0


class FileAnalyzer:
    """PPTX 文件分析器 - 执行 Native Parsing。"""

    def __init__(self, pptx_path: str):
        self.pptx_path = pptx_path
        self.prs = Presentation(pptx_path)
        self.slide_width = self.prs.slide_width
        self.slide_height = self.prs.slide_height

    def analyze_all_slides(self) -> list[SlideData]:
        """分析所有幻灯片。"""
        slides_data = []
        for idx, slide in enumerate(self.prs.slides):
            slide_data = self._analyze_slide(slide, idx)
            slides_data.append(slide_data)
        return slides_data

    def _analyze_slide(self, slide, slide_index: int) -> SlideData:
        """分析单页幻灯片。"""
        data = SlideData(
            slide_index=slide_index,
            width=self.slide_width,
            height=self.slide_height,
        )

        # 提取标题
        data.title = self._extract_title(slide)

        # 遍历所有 shape
        for shape in slide.shapes:
            self._process_shape(shape, data)

        return data

    def _extract_title(self, slide) -> str:
        """提取幻灯片标题。"""
        # 优先用 title placeholder
        for shape in slide.shapes:
            if shape.is_placeholder and shape.placeholder_format.idx == 0:
                if shape.has_text_frame:
                    return self._get_text(shape)
        # 回退：找最大字号的文本
        max_size = 0
        title = ""
        for shape in slide.shapes:
            if shape.has_text_frame:
                size = self._get_max_font_size(shape)
                if size > max_size:
                    max_size = size
                    title = self._get_text(shape)
        return title

    def _process_shape(self, shape, data: SlideData):
        """处理单个 shape。"""
        bbox = self._get_bbox(shape)

        if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
            self._process_table(shape, data, bbox)
        elif shape.shape_type == MSO_SHAPE_TYPE.CHART:
            self._process_chart(shape, data, bbox)
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            self._process_picture(shape, data, bbox)
        elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            # 递归处理组
            for child in shape.shapes:
                self._process_shape(child, data)
        elif shape.has_text_frame:
            text = self._get_text(shape)
            if text:
                font_size = self._get_max_font_size(shape)
                is_title = shape.is_placeholder and shape.placeholder_format.idx == 0
                data.texts.append(TextElement(
                    text=text,
                    bbox=bbox,
                    font_size=font_size,
                    is_title=is_title,
                    shape_type=str(shape.shape_type),
                ))
        elif shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            # 形状（流程图节点等）
            text = self._get_text(shape) if shape.has_text_frame else ""
            data.shapes.append(ShapeElement(
                shape_type=str(shape.shape_type),
                text=text,
                bbox=bbox,
            ))

    def _process_table(self, shape, data: SlideData, bbox):
        """处理表格。"""
        tbl = shape.table
        nrows = len(tbl.rows)
        ncols = len(tbl.columns)
        cells = []
        for r in range(nrows):
            row = []
            for c in range(ncols):
                cell = tbl.cell(r, c)
                row.append(cell.text.strip())
            cells.append(row)

        # 生成 HTML
        html = self._table_to_html(tbl)
        data.tables.append(TableElement(
            rows=nrows,
            cols=ncols,
            cells=cells,
            bbox=bbox,
            html=html,
        ))

    def _table_to_html(self, tbl) -> str:
        """将表格转为 HTML。"""
        nrows = len(tbl.rows)
        ncols = len(tbl.columns)
        html = "<table>\n"
        for r in range(nrows):
            html += "  <tr>\n"
            for c in range(ncols):
                cell = tbl.cell(r, c)
                text = cell.text.strip()
                colspan = getattr(cell, "span_width", 1) or 1
                rowspan = getattr(cell, "span_height", 1) or 1
                attrs = ""
                if colspan > 1:
                    attrs += f' colspan="{colspan}"'
                if rowspan > 1:
                    attrs += f' rowspan="{rowspan}"'
                tag = "th" if r == 0 else "td"
                html += f"    <{tag}{attrs}>{text}</{tag}>\n"
            html += "  </tr>\n"
        html += "</table>"
        return html

    def _process_chart(self, shape, data: SlideData, bbox):
        """处理图表。"""
        chart = shape.chart
        chart_type = str(chart.chart_type)

        # 提取数据
        series_data = []
        for s in chart.series:
            vals = list(s.values)
            name = getattr(s, 'name', f"Series {len(series_data)}")
            series_data.append({"name": name, "values": vals})

        categories = []
        try:
            for pt in chart.plots[0].categories:
                categories.append(str(pt))
        except Exception:
            pass

        data.charts.append(ChartElement(
            chart_type=chart_type,
            data={"series": series_data, "categories": categories},
            bbox=bbox,
        ))

    def _process_picture(self, shape, data: SlideData, bbox):
        """处理图片（保存原始字节供 VLM 处理）。"""
        try:
            image = shape.image
            data.images.append(ImageElement(
                image_bytes=image.blob,
                bbox=bbox,
                shape_type="picture",
            ))
        except Exception:
            pass

    def _get_bbox(self, shape) -> tuple[int, int, int, int]:
        """获取 shape 的边界框 (x1, y1, x2, y2)。"""
        return (shape.left, shape.top, shape.left + shape.width, shape.top + shape.height)

    def _get_text(self, shape) -> str:
        """提取 shape 的全部文本。"""
        if not shape.has_text_frame:
            return ""
        parts = []
        for p in shape.text_frame.paragraphs:
            line = p.text.strip()
            if line:
                parts.append(line)
        return "\n".join(parts)

    def _get_max_font_size(self, shape) -> float:
        """获取 shape 中最大字号。"""
        if not shape.has_text_frame:
            return 0
        sizes = []
        for p in shape.text_frame.paragraphs:
            for run in p.runs:
                if run.font.size:
                    sizes.append(run.font.size.pt)
        return max(sizes) if sizes else 0
